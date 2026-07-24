from __future__ import annotations

import csv
import io
from dataclasses import dataclass
from pathlib import Path

from openpyxl import load_workbook

REQUIRED_COLUMNS = {"merge_group_key", "source_sis_id"}


class IntakeError(ValueError):
    pass


@dataclass(frozen=True)
class IntakeRow:
    merge_group_key: str
    source_sis_id: str
    destination_subaccount: int | None = None


def _normalize_rows(rows: list[dict[str, object]]) -> list[IntakeRow]:
    if not rows:
        raise IntakeError("The intake contains no section rows")
    columns = {str(key).strip().lower() for key in rows[0]}
    missing = REQUIRED_COLUMNS - columns
    if missing:
        raise IntakeError(f"Missing required columns: {', '.join(sorted(missing))}")
    result: list[IntakeRow] = []
    for index, raw in enumerate(rows, start=2):
        normalized = {str(key).strip().lower(): value for key, value in raw.items()}
        group_key = str(normalized.get("merge_group_key") or "").strip()
        source_sis_id = str(normalized.get("source_sis_id") or "").strip()
        if not group_key or not source_sis_id:
            raise IntakeError(f"Row {index} must include merge_group_key and source_sis_id")
        account_value = normalized.get("destination_subaccount")
        try:
            account_id = int(account_value) if account_value not in (None, "") else None
        except (TypeError, ValueError) as exc:
            raise IntakeError(f"Row {index} destination_subaccount must be a Canvas ID") from exc
        result.append(IntakeRow(group_key, source_sis_id, account_id))
    return result


def parse_csv(content: bytes) -> list[IntakeRow]:
    try:
        text = content.decode("utf-8-sig")
    except UnicodeDecodeError as exc:
        raise IntakeError("CSV files must use UTF-8 encoding") from exc
    reader = csv.DictReader(io.StringIO(text))
    return _normalize_rows(list(reader))


def parse_xlsx(content: bytes) -> list[IntakeRow]:
    try:
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception as exc:
        raise IntakeError("Unable to read the XLSX workbook") from exc
    worksheet = workbook.active
    iterator = worksheet.iter_rows(values_only=True)
    try:
        header = next(iterator)
    except StopIteration as exc:
        raise IntakeError("The workbook is empty") from exc
    names = [str(value or "").strip() for value in header]
    rows = [dict(zip(names, values, strict=False)) for values in iterator if any(values)]
    return _normalize_rows(rows)


def parse_upload(filename: str, content: bytes) -> list[IntakeRow]:
    suffix = Path(filename).suffix.lower()
    if suffix == ".csv":
        return parse_csv(content)
    if suffix == ".xlsx":
        return parse_xlsx(content)
    raise IntakeError("Upload a .csv or .xlsx file")


def parse_manual(text: str) -> list[IntakeRow]:
    rows: list[IntakeRow] = []
    for number, raw_line in enumerate(text.splitlines(), start=1):
        line = raw_line.strip()
        if not line:
            continue
        parts = [part.strip() for part in line.split(",")]
        if len(parts) not in (2, 3):
            raise IntakeError(
                f"Manual row {number} must be destination course group, "
                "source SIS section ID, and optional destination Canvas subaccount ID"
            )
        account_id = int(parts[2]) if len(parts) == 3 and parts[2] else None
        rows.append(IntakeRow(parts[0], parts[1], account_id))
    if not rows:
        raise IntakeError("Enter at least one section row or upload a file")
    return rows
